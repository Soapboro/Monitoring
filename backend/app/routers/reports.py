import io
from datetime import date
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, case
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.units import cm

from app.database import get_db
from app.models.grade import Grade, GradeType
from app.models.attendance import Attendance
from app.models.student import Student
from app.models.group import Group
from app.models.teaching_assignment import TeachingAssignment
from app.models.subject import Subject
from app.models.test_session import TestSession, SessionStatus
from app.models.test import Test
from app.models.teacher import Teacher
from app.models.user import User, UserRole
from app.dependencies import require_teacher, get_current_user

router = APIRouter(prefix="/api/reports", tags=["reports"])


async def _teacher_scope(db: AsyncSession, user: User):
    """Возвращает (is_admin, allowed_group_ids | None, allowed_subject_ids | None).
    Для администратора — None означает «без ограничений»."""
    if user.role == UserRole.admin:
        return True, None, None

    teacher = (await db.execute(
        select(Teacher).where(Teacher.user_id == user.id)
    )).scalar_one_or_none()
    if teacher is None:
        raise HTTPException(status_code=403, detail="Профиль преподавателя не найден")

    rows = (await db.execute(
        select(TeachingAssignment.group_id, TeachingAssignment.subject_id)
        .where(TeachingAssignment.teacher_id == teacher.id)
    )).all()
    return False, {r.group_id for r in rows}, {r.subject_id for r in rows}


def _check_scope(is_admin: bool, group_id, subject_id, allowed_groups, allowed_subjects):
    """Бросает 403, если преподаватель запрашивает чужую группу или дисциплину."""
    if is_admin:
        return
    if group_id is not None and group_id not in allowed_groups:
        raise HTTPException(status_code=403, detail="Нет доступа к этой группе")
    if subject_id is not None and subject_id not in allowed_subjects:
        raise HTTPException(status_code=403, detail="Нет доступа к этой дисциплине")


async def _get_grade_data(
    db: AsyncSession,
    group_id: int | None,
    subject_id: int | None,
    acad_year: str | None,
    grade_type: str | None,
    allowed_groups: set | None = None,
    allowed_subjects: set | None = None,
):
    query = (
        select(
            Student.last_name,
            Student.first_name,
            Student.middle_name,
            Group.name.label("group_name"),
            Subject.name.label("subject_name"),
            TeachingAssignment.acad_year,
            TeachingAssignment.semester,
            Grade.grade_type,
            Grade.value,
            Grade.date_recorded,
        )
        .join(Grade, Grade.student_id == Student.id)
        .join(Group, Student.group_id == Group.id)
        .join(TeachingAssignment, Grade.assignment_id == TeachingAssignment.id)
        .join(Subject, TeachingAssignment.subject_id == Subject.id)
        .where(Grade.value.isnot(None))
        .order_by(Group.name, Student.last_name, Student.first_name, Subject.name)
    )
    if group_id:
        query = query.where(TeachingAssignment.group_id == group_id)
    elif allowed_groups is not None:
        query = query.where(TeachingAssignment.group_id.in_(allowed_groups))
    if subject_id:
        query = query.where(TeachingAssignment.subject_id == subject_id)
    elif allowed_subjects is not None:
        query = query.where(TeachingAssignment.subject_id.in_(allowed_subjects))
    if acad_year:
        query = query.where(TeachingAssignment.acad_year == acad_year)
    if grade_type:
        query = query.where(Grade.grade_type == grade_type)

    result = await db.execute(query)
    return result.mappings().all()


@router.get("/grades/excel")
async def grades_excel(
    group_id: int | None = Query(None),
    subject_id: int | None = Query(None),
    acad_year: str | None = Query(None),
    grade_type: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    is_admin, allowed_groups, allowed_subjects = await _teacher_scope(db, current_user)
    _check_scope(is_admin, group_id, subject_id, allowed_groups, allowed_subjects)
    rows = await _get_grade_data(db, group_id, subject_id, acad_year, grade_type, allowed_groups, allowed_subjects)

    wb = Workbook()
    ws = wb.active
    ws.title = "Оценки"

    headers = ["Группа", "Фамилия", "Имя", "Отчество", "Дисциплина", "Уч. год", "Семестр", "Тип оценки", "Оценка", "Дата"]
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_alignment = Alignment(horizontal="center", wrap_text=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, r in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=r["group_name"])
        ws.cell(row=row_idx, column=2, value=r["last_name"])
        ws.cell(row=row_idx, column=3, value=r["first_name"])
        ws.cell(row=row_idx, column=4, value=r["middle_name"] or "")
        ws.cell(row=row_idx, column=5, value=r["subject_name"])
        ws.cell(row=row_idx, column=6, value=r["acad_year"])
        ws.cell(row=row_idx, column=7, value=r["semester"])
        ws.cell(row=row_idx, column=8, value=str(r["grade_type"]))
        ws.cell(row=row_idx, column=9, value=float(r["value"]) if r["value"] is not None else "")
        ws.cell(row=row_idx, column=10, value=str(r["date_recorded"]))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=grades.xlsx"},
    )


@router.get("/grades/pdf")
async def grades_pdf(
    group_id: int | None = Query(None),
    subject_id: int | None = Query(None),
    acad_year: str | None = Query(None),
    grade_type: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    is_admin, allowed_groups, allowed_subjects = await _teacher_scope(db, current_user)
    _check_scope(is_admin, group_id, subject_id, allowed_groups, allowed_subjects)
    rows = await _get_grade_data(db, group_id, subject_id, acad_year, grade_type, allowed_groups, allowed_subjects)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1*cm, rightMargin=1*cm)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontSize=14, spaceAfter=12)
    elements = [Paragraph("Ведомость успеваемости", title_style), Spacer(1, 0.3*cm)]

    table_data = [["Группа", "ФИО", "Дисциплина", "Год", "Сем.", "Тип", "Оценка", "Дата"]]
    for r in rows:
        fio = f"{r['last_name']} {r['first_name']} {r['middle_name'] or ''}".strip()
        table_data.append([
            r["group_name"],
            fio,
            r["subject_name"],
            r["acad_year"],
            str(r["semester"]),
            str(r["grade_type"]),
            str(r["value"]) if r["value"] is not None else "—",
            str(r["date_recorded"]),
        ])

    col_widths = [2.5*cm, 5*cm, 6*cm, 2*cm, 1.5*cm, 2.5*cm, 2*cm, 2.5*cm]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2FF")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("WORDWRAP", (0, 0), (-1, -1), True),
    ]))
    elements.append(t)

    doc.build(elements)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=grades.pdf"},
    )


@router.get("/attendance/excel")
async def attendance_excel(
    group_id: int | None = Query(None),
    subject_id: int | None = Query(None),
    assignment_id: int | None = Query(None),
    date_from: date | None = Query(None),
    date_to: date | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    is_admin, allowed_groups, allowed_subjects = await _teacher_scope(db, current_user)
    _check_scope(is_admin, group_id, subject_id, allowed_groups, allowed_subjects)

    query = (
        select(
            Student.last_name,
            Student.first_name,
            Group.name.label("group_name"),
            Subject.name.label("subject_name"),
            Attendance.lesson_date,
            Attendance.is_present,
            Attendance.comment,
        )
        .join(Attendance, Attendance.student_id == Student.id)
        .join(Group, Student.group_id == Group.id)
        .join(TeachingAssignment, Attendance.assignment_id == TeachingAssignment.id)
        .join(Subject, TeachingAssignment.subject_id == Subject.id)
        .order_by(Group.name, Student.last_name, Attendance.lesson_date)
    )
    if group_id:
        query = query.where(TeachingAssignment.group_id == group_id)
    elif allowed_groups is not None:
        query = query.where(TeachingAssignment.group_id.in_(allowed_groups))
    if subject_id:
        query = query.where(TeachingAssignment.subject_id == subject_id)
    elif allowed_subjects is not None:
        query = query.where(TeachingAssignment.subject_id.in_(allowed_subjects))
    if assignment_id:
        query = query.where(Attendance.assignment_id == assignment_id)
    if date_from:
        query = query.where(Attendance.lesson_date >= date_from)
    if date_to:
        query = query.where(Attendance.lesson_date <= date_to)

    result = await db.execute(query)
    rows = result.mappings().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Посещаемость"

    headers = ["Группа", "Фамилия", "Имя", "Дисциплина", "Дата", "Присутствовал", "Комментарий"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, r in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=r["group_name"])
        ws.cell(row=row_idx, column=2, value=r["last_name"])
        ws.cell(row=row_idx, column=3, value=r["first_name"])
        ws.cell(row=row_idx, column=4, value=r["subject_name"])
        ws.cell(row=row_idx, column=5, value=str(r["lesson_date"]))
        ws.cell(row=row_idx, column=6, value="Да" if r["is_present"] else "Нет")
        ws.cell(row=row_idx, column=7, value=r["comment"] or "")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=attendance.xlsx"},
    )


async def _get_student_report_data(db: AsyncSession, student_id: int):
    student_q = (
        select(Student, Group.name.label("group_name"))
        .join(Group, Student.group_id == Group.id)
        .where(Student.id == student_id)
    )
    row = (await db.execute(student_q)).mappings().one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="Студент не найден")
    student = row["Student"]
    group_name = row["group_name"]

    grades_q = (
        select(
            Subject.name.label("subject"),
            TeachingAssignment.acad_year,
            TeachingAssignment.semester,
            func.count(Grade.id).label("grades_count"),
            func.round(func.avg(Grade.value), 2).label("avg_grade"),
            func.min(Grade.value).label("min_grade"),
            func.max(Grade.value).label("max_grade"),
        )
        .join(TeachingAssignment, Grade.assignment_id == TeachingAssignment.id)
        .join(Subject, TeachingAssignment.subject_id == Subject.id)
        .where(Grade.student_id == student_id, Grade.value.isnot(None))
        .group_by(Subject.name, TeachingAssignment.acad_year, TeachingAssignment.semester)
        .order_by(TeachingAssignment.acad_year, TeachingAssignment.semester, Subject.name)
    )
    grades = (await db.execute(grades_q)).mappings().all()

    attend_q = (
        select(
            Subject.name.label("subject"),
            func.count(Attendance.id).label("total"),
            func.sum(case((Attendance.is_present == True, 1), else_=0)).label("present"),
        )
        .join(TeachingAssignment, Attendance.assignment_id == TeachingAssignment.id)
        .join(Subject, TeachingAssignment.subject_id == Subject.id)
        .where(Attendance.student_id == student_id)
        .group_by(Subject.name)
        .order_by(Subject.name)
    )
    attendance = (await db.execute(attend_q)).mappings().all()

    tests_q = (
        select(
            Test.title,
            Subject.name.label("subject"),
            func.count(TestSession.id).label("attempts"),
            func.round(
                func.max(
                    case(
                        (TestSession.score_max > 0,
                         TestSession.score_total / TestSession.score_max * 100),
                        else_=None,
                    )
                ), 1
            ).label("best_pct"),
            func.max(case((TestSession.passed == True, 1), else_=0)).label("passed"),
        )
        .join(Test, TestSession.test_id == Test.id)
        .join(Subject, Test.subject_id == Subject.id)
        .where(TestSession.student_id == student_id, TestSession.status == SessionStatus.completed)
        .group_by(Test.id, Test.title, Subject.name)
        .order_by(Subject.name, Test.title)
    )
    tests = (await db.execute(tests_q)).mappings().all()

    return student, group_name, grades, attendance, tests


@router.get("/student/excel")
async def student_excel(
    student_id: int = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    is_admin, allowed_groups, _ = await _teacher_scope(db, current_user)
    if not is_admin:
        grp = (await db.execute(select(Student.group_id).where(Student.id == student_id))).scalar_one_or_none()
        if grp is None or grp not in allowed_groups:
            raise HTTPException(status_code=403, detail="Нет доступа к этому студенту")
    student, group_name, grades, attendance, tests = await _get_student_report_data(db, student_id)
    full_name = " ".join(filter(None, [student.last_name, student.first_name, student.middle_name]))

    wb = Workbook()

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    hdr_align = Alignment(horizontal="center", wrap_text=True)

    def write_headers(ws, headers):
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = hdr_align

    def autowidth(ws):
        for col in ws.columns:
            mx = max((len(str(cell.value or "")) for cell in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(mx + 3, 45)

    # Sheet 1 — Grades by subject
    ws1 = wb.active
    ws1.title = "Успеваемость"
    ws1.cell(row=1, column=1, value=f"Студент: {full_name}").font = Font(bold=True)
    ws1.cell(row=2, column=1, value=f"Группа: {group_name}")
    ws1.append([])
    write_headers(ws1, ["Дисциплина", "Уч. год", "Сем.", "Кол-во оценок", "Средний балл", "Мин.", "Макс."])
    ws1.row_dimensions[4].height = 32
    for r in grades:
        ws1.append([
            r["subject"], r["acad_year"], r["semester"],
            r["grades_count"],
            float(r["avg_grade"]) if r["avg_grade"] is not None else "",
            float(r["min_grade"]) if r["min_grade"] is not None else "",
            float(r["max_grade"]) if r["max_grade"] is not None else "",
        ])
    autowidth(ws1)

    # Sheet 2 — Attendance
    ws2 = wb.create_sheet("Посещаемость")
    write_headers(ws2, ["Дисциплина", "Всего занятий", "Присутствовал", "% посещ."])
    for r in attendance:
        total = r["total"] or 0
        present = int(r["present"] or 0)
        rate = round(present / total * 100, 1) if total > 0 else ""
        ws2.append([r["subject"], total, present, rate])
    autowidth(ws2)

    # Sheet 3 — Tests
    ws3 = wb.create_sheet("Тесты")
    write_headers(ws3, ["Тест", "Дисциплина", "Попыток", "Лучший результат %", "Сдан"])
    for r in tests:
        ws3.append([
            r["title"], r["subject"], r["attempts"],
            float(r["best_pct"]) if r["best_pct"] is not None else "",
            "Да" if r["passed"] else "Нет",
        ])
    autowidth(ws3)

    fname = f"student_{student_id}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@router.get("/student/pdf")
async def student_pdf(
    student_id: int = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    is_admin, allowed_groups, _ = await _teacher_scope(db, current_user)
    if not is_admin:
        grp = (await db.execute(select(Student.group_id).where(Student.id == student_id))).scalar_one_or_none()
        if grp is None or grp not in allowed_groups:
            raise HTTPException(status_code=403, detail="Нет доступа к этому студенту")
    student, group_name, grades, attendance, tests = await _get_student_report_data(db, student_id)
    full_name = " ".join(filter(None, [student.last_name, student.first_name, student.middle_name]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=13, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceAfter=4, spaceBefore=10)
    normal = ParagraphStyle("normal", parent=styles["Normal"], fontSize=9, spaceAfter=6)

    base_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2FF")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]

    elems = [
        Paragraph(f"Individual report: {full_name}", h1),
        Paragraph(f"Group: {group_name}", normal),
        Spacer(1, 0.3*cm),
    ]

    # Grades table
    elems.append(Paragraph("Performance by Subject", h2))
    if grades:
        tdata = [["Subject", "Year", "Sem.", "Count", "Avg", "Min", "Max"]]
        for r in grades:
            tdata.append([
                r["subject"], r["acad_year"], str(r["semester"]), str(r["grades_count"]),
                str(r["avg_grade"]), str(r["min_grade"]), str(r["max_grade"]),
            ])
        t = Table(tdata, colWidths=[7*cm, 2.2*cm, 1.4*cm, 2*cm, 1.8*cm, 1.8*cm, 1.8*cm], repeatRows=1)
        t.setStyle(TableStyle(base_style))
        elems.append(t)
    else:
        elems.append(Paragraph("No grades recorded.", normal))

    # Attendance table
    elems.append(Paragraph("Attendance by Subject", h2))
    if attendance:
        tdata = [["Subject", "Total", "Present", "Rate %"]]
        for r in attendance:
            total = r["total"] or 0
            present = int(r["present"] or 0)
            rate = f"{round(present / total * 100, 1)}" if total > 0 else "—"
            tdata.append([r["subject"], str(total), str(present), rate])
        t = Table(tdata, colWidths=[9*cm, 2.5*cm, 2.5*cm, 2.5*cm], repeatRows=1)
        t.setStyle(TableStyle(base_style))
        elems.append(t)
    else:
        elems.append(Paragraph("No attendance records.", normal))

    # Tests table
    elems.append(Paragraph("Test Results", h2))
    if tests:
        tdata = [["Test", "Subject", "Attempts", "Best %", "Passed"]]
        for r in tests:
            tdata.append([
                r["title"], r["subject"], str(r["attempts"]),
                str(r["best_pct"]) if r["best_pct"] is not None else "—",
                "Yes" if r["passed"] else "No",
            ])
        t = Table(tdata, colWidths=[7*cm, 5*cm, 2.2*cm, 2.2*cm, 2*cm], repeatRows=1)
        t.setStyle(TableStyle(base_style))
        elems.append(t)
    else:
        elems.append(Paragraph("No test sessions.", normal))

    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=student_{student_id}.pdf"},
    )


@router.get("/analytics/excel")
async def analytics_excel(
    group_id: int | None = Query(None),
    subject_id: int | None = Query(None),
    acad_year: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    is_admin, allowed_groups, allowed_subjects = await _teacher_scope(db, current_user)
    _check_scope(is_admin, group_id, subject_id, allowed_groups, allowed_subjects)
    """Агрегированный аналитический отчёт: средние баллы по группам, дисциплинам, рейтинг студентов."""
    wb = Workbook()

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    hdr_align = Alignment(horizontal="center", wrap_text=True)

    def write_headers(ws, headers, start_row=1):
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=start_row, column=col, value=h)
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = hdr_align

    def autowidth(ws):
        for col in ws.columns:
            mx = max((len(str(cell.value or "")) for cell in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(mx + 3, 45)

    # Sheet 1 — Group averages by subject
    ws1 = wb.active
    ws1.title = "Ср. балл по дисциплинам"
    group_subj_q = (
        select(
            Group.name.label("group"),
            Subject.name.label("subject"),
            TeachingAssignment.acad_year,
            TeachingAssignment.semester,
            func.count(func.distinct(Grade.student_id)).label("students"),
            func.round(func.avg(Grade.value), 2).label("avg_grade"),
            func.min(Grade.value).label("min_grade"),
            func.max(Grade.value).label("max_grade"),
            func.count(Grade.id).label("grades_count"),
        )
        .join(TeachingAssignment, Grade.assignment_id == TeachingAssignment.id)
        .join(Group, TeachingAssignment.group_id == Group.id)
        .join(Subject, TeachingAssignment.subject_id == Subject.id)
        .where(Grade.value.isnot(None))
        .group_by(Group.name, Subject.name, TeachingAssignment.acad_year, TeachingAssignment.semester)
        .order_by(TeachingAssignment.acad_year, Group.name, Subject.name)
    )
    if group_id:
        group_subj_q = group_subj_q.where(TeachingAssignment.group_id == group_id)
    elif allowed_groups is not None:
        group_subj_q = group_subj_q.where(TeachingAssignment.group_id.in_(allowed_groups))
    if subject_id:
        group_subj_q = group_subj_q.where(TeachingAssignment.subject_id == subject_id)
    elif allowed_subjects is not None:
        group_subj_q = group_subj_q.where(TeachingAssignment.subject_id.in_(allowed_subjects))
    if acad_year:
        group_subj_q = group_subj_q.where(TeachingAssignment.acad_year == acad_year)
    rows1 = (await db.execute(group_subj_q)).mappings().all()
    write_headers(ws1, ["Группа", "Дисциплина", "Уч. год", "Сем.", "Студентов", "Ср. балл", "Мин.", "Макс.", "Кол-во оценок"])
    for r in rows1:
        ws1.append([
            r["group"], r["subject"], r["acad_year"], r["semester"],
            r["students"],
            float(r["avg_grade"]) if r["avg_grade"] is not None else "",
            float(r["min_grade"]) if r["min_grade"] is not None else "",
            float(r["max_grade"]) if r["max_grade"] is not None else "",
            r["grades_count"],
        ])
    autowidth(ws1)

    # Sheet 2 — Student rankings
    ws2 = wb.create_sheet("Рейтинг студентов")
    student_q = (
        select(
            (Student.last_name + " " + Student.first_name).label("name"),
            Group.name.label("group"),
            func.round(func.avg(Grade.value), 2).label("avg_grade"),
            func.count(Grade.id).label("grades_count"),
        )
        .join(Grade, Grade.student_id == Student.id)
        .join(Group, Student.group_id == Group.id)
        .join(TeachingAssignment, Grade.assignment_id == TeachingAssignment.id)
        .where(Grade.value.isnot(None))
        .group_by(Student.id, Student.last_name, Student.first_name, Group.name)
        .order_by(func.avg(Grade.value).desc())
    )
    if group_id:
        student_q = student_q.where(TeachingAssignment.group_id == group_id)
    elif allowed_groups is not None:
        student_q = student_q.where(TeachingAssignment.group_id.in_(allowed_groups))
    if acad_year:
        student_q = student_q.where(TeachingAssignment.acad_year == acad_year)
    rows2 = (await db.execute(student_q)).mappings().all()
    write_headers(ws2, ["№", "ФИО", "Группа", "Ср. балл", "Кол-во оценок"])
    for i, r in enumerate(rows2, 1):
        ws2.append([i, r["name"], r["group"], float(r["avg_grade"]) if r["avg_grade"] is not None else "", r["grades_count"]])
    autowidth(ws2)

    # Sheet 3 — Attendance summary
    ws3 = wb.create_sheet("Посещаемость")
    att_q = (
        select(
            Group.name.label("group"),
            Subject.name.label("subject"),
            func.count(Attendance.id).label("total"),
            func.sum(case((Attendance.is_present == True, 1), else_=0)).label("present"),
        )
        .join(TeachingAssignment, Attendance.assignment_id == TeachingAssignment.id)
        .join(Group, TeachingAssignment.group_id == Group.id)
        .join(Subject, TeachingAssignment.subject_id == Subject.id)
        .group_by(Group.name, Subject.name)
        .order_by(Group.name, Subject.name)
    )
    if group_id:
        att_q = att_q.where(TeachingAssignment.group_id == group_id)
    elif allowed_groups is not None:
        att_q = att_q.where(TeachingAssignment.group_id.in_(allowed_groups))
    if subject_id:
        att_q = att_q.where(TeachingAssignment.subject_id == subject_id)
    elif allowed_subjects is not None:
        att_q = att_q.where(TeachingAssignment.subject_id.in_(allowed_subjects))
    rows3 = (await db.execute(att_q)).mappings().all()
    write_headers(ws3, ["Группа", "Дисциплина", "Всего", "Присутствовало", "% посещ."])
    for r in rows3:
        total = r["total"] or 0
        present = int(r["present"] or 0)
        rate = round(present / total * 100, 1) if total > 0 else ""
        ws3.append([r["group"], r["subject"], total, present, rate])
    autowidth(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=analytics.xlsx"},
    )
